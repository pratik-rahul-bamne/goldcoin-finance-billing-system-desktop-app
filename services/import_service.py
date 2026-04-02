"""Validate and import .xlsx backups into SQLite (append or replace business data)."""

from __future__ import annotations

import io
import sqlite3
from typing import Any, BinaryIO

from openpyxl import Workbook, load_workbook

from services.backup_service import TABLES, _safe_sheet_title

# Tables restored from file on replace (settings handled separately via flag).
_DATA_TABLES: tuple[str, ...] = (
    "service_catalog",
    "customers",
    "services",
    "payments",
)


def _pragma_columns(conn: sqlite3.Connection, table: str) -> list[str]:
    cur = conn.execute(f'PRAGMA table_info("{table}")')
    return [row[1] for row in cur.fetchall()]


def _normalize_cell(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, str) and v.strip() == "":
        return ""
    return v


def _row_to_dict(headers: list[str], row: tuple[Any, ...]) -> dict[str, Any]:
    return {h: _normalize_cell(row[i] if i < len(row) else None) for i, h in enumerate(headers)}


def read_sheet_as_dicts(ws) -> tuple[list[str], list[dict[str, Any]]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"Sheet {ws.title!r} is empty (expected header row).")
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    if not headers or all(h == "" for h in headers):
        raise ValueError(f"Sheet {ws.title!r} has no column headers.")
    data = []
    for r in rows[1:]:
        if r is None or all(c is None or str(c).strip() == "" for c in r):
            continue
        padded = tuple(r[i] if i < len(r) else None for i in range(len(headers)))
        data.append(_row_to_dict(headers, padded))
    return headers, data


def _workbook_from_stream(stream: BinaryIO) -> Workbook:
    bio = io.BytesIO(stream.read())
    bio.seek(0)
    return load_workbook(bio, read_only=False, data_only=True)


def _find_sheet(wb: Workbook, table: str):
    target = _safe_sheet_title(table)
    for name in wb.sheetnames:
        if name == table or name == target:
            return wb[name]
    raise ValueError(f"Missing sheet for table {table!r} (expected title {target!r}).")


def validate_workbook_stream(stream: BinaryIO, conn: sqlite3.Connection) -> dict[str, Any]:
    """
    Load workbook and verify each TABLE has a sheet whose header row matches the DB schema.
    Returns metadata dict suitable for preview (counts per table).
    """
    wb = _workbook_from_stream(stream)
    summary: dict[str, Any] = {"tables": {}, "ok": True}
    for table in TABLES:
        expected = _pragma_columns(conn, table)
        ws = _find_sheet(wb, table)
        headers, data = read_sheet_as_dicts(ws)
        missing = [c for c in expected if c not in headers]
        extra = [c for c in headers if c not in expected]
        if missing:
            raise ValueError(
                f"Sheet {table!r}: missing columns {missing}. Expected {expected}, got {headers}."
            )
        if extra:
            # Allow extra columns in file but ignore on import
            pass
        summary["tables"][table] = {
            "row_count": len(data),
            "columns_ok": expected,
        }
    wb.close()
    return summary


def import_workbook_stream(
    stream: BinaryIO,
    conn: sqlite3.Connection,
    *,
    mode: str,
    merge_settings: bool,
    duplicate_mobile: str = "reuse",
) -> dict[str, Any]:
    """
    mode: 'replace' — delete service_catalog, customers, services, payments then import those sheets.
           'append' — merge rows; customers keyed by backup id with mobile dedup.
    merge_settings: if True, INSERT OR REPLACE settings keys from file (can change admin password).
    duplicate_mobile: 'reuse' maps old customer id to existing row with same mobile; 'error' raises.
    """
    if mode not in ("append", "replace"):
        raise ValueError("mode must be 'append' or 'replace'.")
    if duplicate_mobile not in ("reuse", "error"):
        raise ValueError("duplicate_mobile must be 'reuse' or 'error'.")

    wb = _workbook_from_stream(stream)
    try:
        # Validate all sheets first
        cache: dict[str, tuple[list[str], list[dict[str, Any]]]] = {}
        for table in TABLES:
            expected = _pragma_columns(conn, table)
            ws = _find_sheet(wb, table)
            headers, data = read_sheet_as_dicts(ws)
            missing = [c for c in expected if c not in headers]
            if missing:
                raise ValueError(
                    f"Sheet {table!r}: missing columns {missing}. Expected {expected}."
                )
            cache[table] = (expected, data)

        stats = {"customers_inserted": 0, "customers_reused": 0, "settings_keys": 0}

        cur = conn.cursor()
        try:
            if mode == "replace":
                cur.execute("DELETE FROM payments")
                cur.execute("DELETE FROM services")
                cur.execute("DELETE FROM customers")
                cur.execute("DELETE FROM service_catalog")

            id_map: dict[int, int] = {}

            if merge_settings:
                _, srows = cache["settings"]
                for row in srows:
                    key = row.get("key")
                    val = row.get("value")
                    if key is None:
                        continue
                    cur.execute(
                        "INSERT OR REPLACE INTO settings (key, value) VALUES (?,?)",
                        (str(key), str(val) if val is not None else ""),
                    )
                    stats["settings_keys"] += 1

            # service_catalog
            cols, rows = cache["service_catalog"]
            insert_cols = [c for c in cols if c != "id"]
            placeholders = ",".join(["?"] * len(insert_cols))
            col_list = ",".join(f'"{c}"' for c in insert_cols)
            for row in rows:
                vals = []
                for c in insert_cols:
                    v = row.get(c)
                    if c == "is_active" and v is not None:
                        try:
                            v = int(v)
                        except (TypeError, ValueError):
                            v = 1
                    vals.append(v)
                if mode == "replace":
                    cur.execute(
                        f'INSERT INTO service_catalog ({col_list}) VALUES ({placeholders})',
                        vals,
                    )
                else:
                    cur.execute(
                        f"""INSERT OR IGNORE INTO service_catalog ({col_list}) VALUES ({placeholders})""",
                        vals,
                    )

            # customers
            cols, crows = cache["customers"]
            c_insert_cols = [c for c in cols if c != "id"]
            c_col_list = ",".join(f'"{c}"' for c in c_insert_cols)
            c_ph = ",".join(["?"] * len(c_insert_cols))

            for row in crows:
                old_id = row.get("id")
                try:
                    old_id_int = int(old_id) if old_id is not None else None
                except (TypeError, ValueError):
                    old_id_int = None
                if old_id_int is None:
                    raise ValueError("customers sheet: each row needs a valid integer id from backup.")

                mobile = str(row.get("mobile") or "").strip()
                if mode == "append" and mobile:
                    ex = cur.execute(
                        "SELECT id FROM customers WHERE mobile = ?",
                        (mobile,),
                    ).fetchone()
                    if ex:
                        if duplicate_mobile == "error":
                            raise ValueError(f"Duplicate mobile on append: {mobile!r}")
                        id_map[old_id_int] = int(ex[0])
                        stats["customers_reused"] += 1
                        continue

                vals = [row.get(c) for c in c_insert_cols]
                cur.execute(
                    f'INSERT INTO customers ({c_col_list}) VALUES ({c_ph})',
                    vals,
                )
                id_map[old_id_int] = int(cur.lastrowid)
                stats["customers_inserted"] += 1

            # services
            cols, srows = cache["services"]
            s_insert_cols = [c for c in cols if c != "id"]
            s_col_list = ",".join(f'"{c}"' for c in s_insert_cols)
            s_ph = ",".join(["?"] * len(s_insert_cols))
            for row in srows:
                old_cid = row.get("customer_id")
                try:
                    oc = int(old_cid)
                except (TypeError, ValueError):
                    raise ValueError("services sheet: customer_id must be integer.")
                new_cid = id_map.get(oc)
                if new_cid is None:
                    raise ValueError(
                        f"services sheet: customer_id {oc} not in imported customer id map."
                    )
                vals = []
                for c in s_insert_cols:
                    if c == "customer_id":
                        vals.append(new_cid)
                    else:
                        vals.append(row.get(c))
                cur.execute(
                    f'INSERT INTO services ({s_col_list}) VALUES ({s_ph})',
                    vals,
                )

            # payments
            cols, prows = cache["payments"]
            p_insert_cols = [c for c in cols if c != "id"]
            p_col_list = ",".join(f'"{c}"' for c in p_insert_cols)
            p_ph = ",".join(["?"] * len(p_insert_cols))
            for row in prows:
                old_cid = row.get("customer_id")
                try:
                    oc = int(old_cid)
                except (TypeError, ValueError):
                    raise ValueError("payments sheet: customer_id must be integer.")
                new_cid = id_map.get(oc)
                if new_cid is None:
                    raise ValueError(
                        f"payments sheet: customer_id {oc} not in imported customer id map."
                    )
                vals = []
                for c in p_insert_cols:
                    if c == "customer_id":
                        vals.append(new_cid)
                    else:
                        vals.append(row.get(c))
                cur.execute(
                    f'INSERT INTO payments ({p_col_list}) VALUES ({p_ph})',
                    vals,
                )

            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            cur.close()

        stats["mode"] = mode
        stats["tables"] = list(_DATA_TABLES if mode == "replace" else TABLES)
        return stats
    finally:
        wb.close()
